"""
試算レポート（PPTX）の雛形を生成するスクリプト。

元のレポート例（doc/20261002 試算レポート例.pptx）を読み込み、
システムで差し替える箇所を {{トークン}} に置き換えた雛形を public/templates/report-template.pptx へ出力する。
実行時の差し替えは lib/export-report-pptx.ts が行う（トークン置換・表の行複製・グラフデータ更新・地図画像差し替え）。

使い方:
  python scripts/build-report-template.py [元pptxのパス]

前提: python-pptx / XlsxWriter（グラフの埋め込みExcel更新に使用）
"""
import copy
import sys
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "doc" / "20261002 試算レポート例.pptx"
DST = ROOT / "public" / "templates" / "report-template.pptx"

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def qn(tag: str) -> str:
    prefix, local = tag.split(":")
    return f"{{{NS_A}}}{local}" if prefix == "a" else tag


def set_runs(paragraph, texts):
    """段落の書式（先頭runのrPr）を保ったまま、run を texts の数だけ作り直す。"""
    p = paragraph._p
    runs = p.findall(qn("a:r"))
    flds = p.findall(qn("a:fld"))
    proto = runs[0] if runs else None
    if proto is None and flds:
        # 日付フィールド等: フィールドの rPr を流用して通常の run にする
        proto = p.makeelement(qn("a:r"), {})
        rpr = flds[0].find(qn("a:rPr"))
        if rpr is not None:
            proto.append(copy.deepcopy(rpr))
        t = p.makeelement(qn("a:t"), {})
        proto.append(t)
    if proto is None:
        proto = p.makeelement(qn("a:r"), {})
        proto.append(p.makeelement(qn("a:t"), {}))
    for el in list(p):
        if el.tag in (qn("a:r"), qn("a:fld"), qn("a:br")):
            p.remove(el)
    end = p.find(qn("a:endParaRPr"))
    for text in texts:
        r = copy.deepcopy(proto)
        # ハイパーリンクは引き継がない（必要な箇所は個別に付け直す）
        rpr = r.find(qn("a:rPr"))
        if rpr is not None:
            for h in rpr.findall(qn("a:hlinkClick")):
                rpr.remove(h)
        r.find(qn("a:t")).text = text
        if end is not None:
            end.addprevious(r)
        else:
            p.append(r)


def set_cell(cell, text):
    """セルの先頭段落だけ残して text を設定する。"""
    tf = cell.text_frame
    paras = tf.paragraphs
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    set_runs(paras[0], [text])


def shape_by_id(slide, shape_id):
    for sh in slide.shapes:
        if sh.shape_id == shape_id:
            return sh
    raise KeyError(f"shape id {shape_id} not found")


def remove_shape(shape):
    el = shape._element
    el.getparent().remove(el)


def delete_rows(table, indexes):
    tbl = table._tbl
    trs = tbl.findall(qn("a:tr"))
    for i in sorted(indexes, reverse=True):
        tbl.remove(trs[i])


def table_by_header(slide, header_text):
    for sh in slide.shapes:
        if sh.has_table and sh.table.cell(0, 0).text.strip() == header_text:
            return sh.table
    raise KeyError(header_text)


def main():
    prs = Presentation(str(SRC))
    s = prs.slides

    # ── 1. 表紙 ──
    cover = s[0]
    for sh in cover.shapes:
        if not sh.has_text_frame:
            continue
        text = sh.text_frame.text
        if "物件住所" in text:
            ps = sh.text_frame.paragraphs
            set_runs(ps[0], ["物件住所：", "{{address}}"])
            set_runs(ps[1], ["坪数：", "{{tsubo}}", "坪"])
            set_runs(ps[2], ["賃料：", "{{rentMan}}", "万円（坪単価約", "{{rentPerTsubo}}", "円）"])
            set_runs(ps[3], ["ロイヤリティ：", "{{royalty}}", "％"])
            set_runs(ps[4], ["グーグルマップ：", "リンク"])
            ps[4].runs[1].hyperlink.address = "https://www.google.com/maps/search/?api=1&query={{mapQuery}}"
        elif sh.text_frame._txBody.find(f".//{qn('a:fld')}") is not None:
            # 作成日（元は自動更新の日付フィールド。開くたびに日付が変わらないよう固定値にする）
            set_runs(sh.text_frame.paragraphs[0], ["{{date}}"])

    # ── 3. 競合情報: サンプルの吹き出し・矢印・赤枠・アイコンを除去（地図画像と見出しだけ残す）──
    comp = s[2]
    for sh in list(comp.shapes):
        if sh.name.startswith("Google Shape;"):
            remove_shape(sh)

    # ── 4. 人口情報: 比較店舗の行・系列を除去し、本物件のみにする ──
    pop = s[3]
    tbl = table_by_header(pop, "有効人口")
    delete_rows(tbl, [2, 3])
    set_cell(tbl.cell(1, 0), "本物件")
    for c, key in enumerate(["pop1", "pop3", "pop5"], start=1):
        set_cell(tbl.cell(1, c), "{{" + key + "}}")
    for sh in pop.shapes:
        if not sh.has_chart:
            continue
        chart = sh.chart
        cd = CategoryChartData(number_format="#,##0")
        if chart.chart_type is not None and "COLUMN" in str(chart.chart_type):
            cd.categories = ["1km", "3km", "5km"]
            cd.add_series("本物件", (0, 0, 0))
        else:
            cats = list(chart.plots[0].categories)
            cd.categories = cats
            cd.add_series("1km圏人口", tuple(0 for _ in cats))
        chart.replace_data(cd)

    # ── 5. 投資金額 / 運営コスト ──
    cost = s[4]
    inv = table_by_header(cost, "内容")
    n = len(inv.rows)
    delete_rows(inv, list(range(2, n - 1)))
    for c, key in enumerate(["inv.label", "inv.amount", "inv.years", "inv.dep"]):
        set_cell(inv.cell(1, c), "{{" + key + "}}")
    set_cell(inv.cell(2, 1), "{{invTotal}}")
    set_cell(inv.cell(2, 3), "{{depTotal}}")

    rc = table_by_header(cost, "月額")
    n = len(rc.rows)
    delete_rows(rc, list(range(2, n - 1)))
    for c, key in enumerate(["rc.label", "rc.unit", "rc.qty", "rc.unitLabel", "rc.amount"]):
        set_cell(rc.cell(1, c), "{{" + key + "}}")
    set_cell(rc.cell(2, 4), "{{rcTotal}}")

    for sh in cost.shapes:
        if not sh.has_text_frame:
            continue
        ps = sh.text_frame.paragraphs
        if "変動費" in sh.text_frame.text:
            set_runs(ps[1], ["・広告費　1～2年目：", "{{adYear2}}", "万円/月　3年目以降", "{{adYear3}}", "万円/月"])
            set_runs(ps[2], ["・変動費　約", "{{variableCost}}", "円/1ユーザー"])
        elif "内装、看板費" in sh.text_frame.text:
            for p in ps:
                if "内装" in "".join(r.text for r in p.runs):
                    set_runs(p, ["・内装、看板費：", "{{interiorPerTsubo}}", "万円/坪にて算出。"])
                elif "マシン" in "".join(r.text for r in p.runs):
                    set_runs(p, ["・マシン購入費：", "{{machinePerTsubo}}", "万円/坪にて算出。"])

    # ── 6〜8. 財務シミュレーション（10期表・KPI）。グラフは実行時に差し替える ──
    for slide in (s[5], s[6], s[7]):
        fin = table_by_header(slide, "")
        for r in range(1, 11):
            for c, key in enumerate(["m", "rev", "gr", "cost", "pt", "pr"], start=1):
                set_cell(fin.cell(r, c), "{{y" + str(r) + "." + key + "}}")
        set_cell(fin.cell(11, 2), "{{tot.rev}}")
        set_cell(fin.cell(11, 4), "{{tot.cost}}")
        set_cell(fin.cell(11, 5), "{{tot.pt}}")
        kpi = table_by_header(slide, "平均単価")
        for r, key in enumerate(["kpi.price", "kpi.be", "kpi.max", "kpi.cc", "kpi.park"]):
            set_cell(kpi.cell(r, 1), "{{" + key + "}}")

    # 差し替え・削除で参照されなくなった関係（サンプルのリンク・画像等）を除去する
    for slide in prs.slides:
        xml = slide._element.xml
        for rid, rel in list(slide.part.rels.items()):
            if rel.reltype.endswith(("/hyperlink", "/image")) and f'"{rid}"' not in xml:
                slide.part.drop_rel(rid)

    prs.core_properties.title = "試算レポート"
    DST.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(DST))
    print(f"wrote {DST} ({DST.stat().st_size:,} bytes) at {date.today()}")


if __name__ == "__main__":
    main()
